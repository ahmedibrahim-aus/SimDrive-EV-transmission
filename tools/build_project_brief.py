"""Build the student project brief, student_package/EV_Gearbox_Project_Brief.docx.

Maintainer tool, not part of the course. Requires python-docx:

    python -m pip install python-docx
    python tools/build_project_brief.py

Paths are resolved relative to the repository root, so it runs from anywhere.
The two figures it embeds are the drive-unit render in tools/ and the three-case
overview that the load-case run writes into the student hand-out.
"""

import os
import re
import zipfile
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "student_package", "EV_Gearbox_Project_Brief.docx")
DRIVE_UNIT_FIGURE = os.path.join(REPO, "tools", "project_brief_drive_unit.png")
OVERVIEW_FIGURE = os.path.join(
    REPO, "student_package", "exports_design_ready", "Overview_AllCases.png")

AUTHOR = "Ahmed Hanafy Ibrahim"
TITLE = "SimDrive: EV Transmission Design"

ACCENT = RGBColor(0x0B, 0x5C, 0x8C)
GRAY = RGBColor(0x55, 0x5A, 0x60)

for figure in (DRIVE_UNIT_FIGURE, OVERVIEW_FIGURE):
    if not os.path.exists(figure):
        raise SystemExit(
            "missing figure: %s\nRun the load cases and buildStudentPackage first."
            % figure)

doc = Document()

# base style
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(8)
st.paragraph_format.line_spacing = 1.15

for lvl, size in ((1, 18), (2, 13.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.color.rgb = ACCENT
    h.font.bold = True
    h.paragraph_format.space_before = Pt(16 if lvl == 1 else 12)
    h.paragraph_format.space_after = Pt(6)


def para(text, italic=False, bold=False, color=None, size=None, align=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic, r.bold = italic, bold
    if color is not None:
        r.font.color.rgb = color
    if size is not None:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def numbered(text):
    doc.add_paragraph(text, style="List Number")


def figure(path, caption):
    doc.add_picture(path, width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, italic=True, color=GRAY, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


# ---------------------------------------------------------------- title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run(TITLE)
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = ACCENT
t.paragraph_format.space_after = Pt(2)

para("Designing the Gearbox of an Electric Car  |  Project Brief",
     color=GRAY, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Third- and fourth-year design project  |  Methods from Shigley's "
     "Mechanical Engineering Design, 11th edition",
     color=GRAY, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Runs alongside a Shigley-based Machine Design course",
     color=GRAY, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

# ---------------------------------------------------------------- hook
doc.add_heading("Start here", level=1)
para("You are going to design the transmission that takes up to 350 N·m (in the "
     "default car) from an electric motor, multiplies it about nine times and "
     "delivers it to the rear wheels, and you will check its conceptual design "
     "against a 150,000 kilometer target. "
     "The life is a classroom estimate, not a service-life prediction.")
para("The loads come from a simulation of the car being driven, not from a "
     "problem statement.")

# ---------------------------------------------------------------- why different
doc.add_heading("What the project adds", level=1)
para("Every stress problem in your design course began the same way. The load "
     "was in the question. Somebody had already decided that the shaft carries "
     "4.2 kN, and your work started after that decision had been made.")
para("This project adds the step before that. The load is not given: you have "
     "to find out what it is, and finding out means reading the results of a "
     "model of the system the component lives in, then arguing about which number out of that model "
     "belongs in which equation.")
para("So the work runs in three stages, and only the middle one looks like the "
     "problems you have solved before.")

doc.add_heading("1. System-level simulation", level=2)
para("A Simulink model of the whole vehicle is driven through three situations: "
     "a city driving cycle, a long climb, and a full-throttle run. The model "
     "accounts for vehicle mass, aerodynamic drag, rolling resistance, the motor "
     "torque and power limits, and a controller that tracks the target speed. "
     "It reports what the transmission actually experiences, second by second.")
para("Your instructor runs this and gives you the results. You do not have to "
     "build the vehicle model, but you do have to understand what it is telling "
     "you, because the next stage depends entirely on reading it correctly.")

doc.add_heading("2. Analysis", level=2)
para("Now the Shigley work begins, and it is the mechanics you already know: "
     "AGMA bending and contact stress on the gear teeth, endurance limits and "
     "Marin factors, stress concentration at every shoulder, the modified "
     "Goodman criterion, rolling-contact bearing life. Chapters 3, 5, 6, 7, 11, "
     "13 and 14.")
para("What has changed is where the numbers come from. The torque in your "
     "Goodman calculation is not given to you. It came out of a city cycle you "
     "can look at and plot. Nor are the rating factors given. Every AGMA factor, "
     "every Marin factor, the material and its grade: you choose each one, "
     "name the figure or table it came from, and defend it. The 150,000 km "
     "service target is fixed; the cycle count of each member over that "
     "target is yours to derive. "
     "An unsourced number costs more points than an arithmetic slip.")

doc.add_heading("3. Optimization", level=2)
para("A design that merely passes is not finished. A shaft thick enough to "
     "pass is easy to find; your task is to find the lightest shaft that still meets "
     "every safety factor at every critical section, which means searching the "
     "feasible geometry rather than picking one and checking it.")
para("Here you find out which constraint controls the design. It is often "
     "not the one you expected, and the components are coupled, so relieving one constraint "
     "usually tightens another.")

doc.add_page_break()

# ---------------------------------------------------------------- the machine
doc.add_heading("The machine", level=1)
para("A rear-wheel-drive passenger electric vehicle. One motor "
     "on the rear axle, feeding a single-speed, two-stage reduction gearbox, "
     "a differential, and a half-shaft to each wheel.")

rows = [
    ("Vehicle mass", "1610 kg"),
    ("Motor peak torque", "350 N·m"),
    ("Motor peak power", "210 kW"),
    ("Gear reduction", "9:1 single speed, in two stages"),
    ("Stage 1 (primary)", "20-tooth pinion, 60-tooth gear, helical"),
    ("Stage 2 (final drive)", "20-tooth pinion, 60-tooth gear, helical"),
    ("Modules and helix angles", "As prescribed in Design_Summary.mat"),
    ("Wheel radius", "0.334 m"),
    ("Target life (gears and bearings)", "150,000 km"),
    ("Safety-factor targets", "Gear bending 1.5, gear contact 1.2; shaft fatigue and yield 1.5"),
]
DEFAULTS_NOTE = ("These are the default values. If your instructor has changed the car or the "
                 "gear train, the values you design with are the ones in Design_Summary.mat.")
tbl = doc.add_table(rows=0, cols=2)
tbl.style = "Light Grid Accent 1"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in rows:
    c = tbl.add_row().cells
    c[0].text = k
    c[1].text = v
    c[0].paragraphs[0].runs[0].bold = True
    for cell in c:
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
para(DEFAULTS_NOTE)

figure(DRIVE_UNIT_FIGURE,
       "The rear drive unit. The motor drives the input pinion; stage 1 turns "
       "the intermediate shaft; the stage-2 pinion on that shaft drives the "
       "output gear, which carries the differential. You design both gear "
       "stages and the intermediate shaft that carries them.")

para("Why two stages. Reaching 9:1 in a single mesh would need a driven gear "
     "larger than the road wheel. Splitting it into two meshes of about 3:1 "
     "keeps every gear small enough to fit beneath the car, and it is how "
     "production EV drive units are built. It also puts "
     "one shaft in the middle that is loaded by both meshes at once, and that "
     "shaft is the one you design.",
     italic=True)

doc.add_page_break()

# ---------------------------------------------------------------- load cases
doc.add_heading("The three load cases", level=1)
para("A gearbox does not see one load. It sees a large torque a few times and a "
     "much smaller one several million times, and those two situations destroy "
     "a component by different mechanisms.")

figure(OVERVIEW_FIGURE,
       "Overview of the three load cases: simulated motor torque and vehicle speed. "
       "For fatigue design use T_mean_eq and T_alt_eq from FTP75_Outputs.mat, not "
       "the raw FTP-75 trace shown here. The data files are in exports_design_ready.")

doc.add_heading("FTP-75 city cycle", level=2)
para("An 1874-second certification cycle of city driving: about two dozen stops "
     "and restarts and hundreds of smaller speed changes. No single load here is dangerous. "
     "Added together over the life of the car they are what can break a shaft by "
     "fatigue, at a stress well below the yield strength. Whether they do on this "
     "car, or whether the single full-throttle peak governs instead, is for you "
     "to show. This cycle also gives the bearings their design load: "
     "T_cubic_mean, the constant torque that Shigley Eq. (11-17) makes "
     "equivalent to this varying duty, from which you compute the L10 life.")

doc.add_heading("Sustained hill climb", level=2)
para("Full throttle on a steep grade (20 degrees by default) until the car settles at the fastest "
     "speed it can hold, with the motor on its power limit. That steady state "
     "has a closed form, so this is the case you use to check the simulation "
     "by hand before you trust any torque from it. As an extension, it also "
     "asks how long your shaft would survive a sustained full-power climb.")

doc.add_heading("Full-throttle acceleration", level=2)
para("Full motor torque from rest. This produces the largest torque the "
     "transmission will ever see. A gear tooth either survives the worst single "
     "load or it fractures, so this is the case for tooth stress and for the "
     "static yield check on the shafts.")

doc.add_page_break()

# ---------------------------------------------------------------- what you design
doc.add_heading("What you design", level=1)
para("Three components, in this order, each one constraining the next.")
numbered("Both helical gear stages. Face width, material and hardness, and "
         "every AGMA rating factor, checked against both bending and pitting "
         "for all four gears.")
numbered("The intermediate shaft. Stepped, with bearing journals and one seat "
         "carrying both gears, checked for fatigue and static yield at every "
         "critical section. Two meshes load it at once; work out how their "
         "forces and thrust couples combine before you calculate.")
numbered("Its two bearings, selected from a real manufacturer's catalog "
         "against rating life, with the thrust factors taken from Shigley "
         "Table 11-1 rather than copied off a catalog page.")
para("The input shaft and its bearings are an optional extension.")

para("These are not three separate problems. The bearing's nominal bore must match "
     "its journal. The journal sets the shoulder geometry. The shoulder "
     "sets the stress concentration that decides the shaft diameter. Choosing a "
     "bigger bearing to feel safe makes your shaft heavier. You will go around  "
     "that loop more than once, and the report asks you to say which constraint "
     "ran out first.")
para("The gear-seat shoulders are not automatically bearing abutments. "
     "Check the catalog abutment diameter and corner radius of each bearing "
     "against its shoulder, and describe the locating spacer or shoulder "
     "separately. Fits and axial retention are outside this project; state in "
     "your report what would have to be detailed before manufacture.")

doc.add_heading("What you hand in", level=1)
numbered("The four completed MATLAB templates.")
numbered("A design report, following the supplied template.")
numbered("Result tables for the gears, shafts and bearings, with units.")
numbered("A scaled drawing of the intermediate shaft, showing diameters, lengths and fillet radii.")
numbered("At least one documented design iteration: what you changed, why, and what moved.")
numbered("The catalog page you took your bearing data from.")

doc.add_heading("How it is graded", level=1)
# Same areas and weights as docs/ASSESSMENT_RUBRIC.md.
mk = [
    ("Load cases and torque handoff (15)", "The right torque for each failure mode, and the hill-climb hand check of the simulation"),
    ("Gear design (25)", "Both stages rated in bending and contact, every factor with its Shigley source, the governing mode compared on a like basis"),
    ("Shaft design (25)", "Reactions in two planes with both thrust couples, DE-Goodman and yield at every critical section"),
    ("Bearing selection (15)", "Equivalent load, Table 11-1 factors, L10 life, a rejected candidate, the catalog cited"),
    ("Integration and iteration (15)", "Compatible bores and seats, every target met, a practical mass, changes justified"),
    ("Engineering communication (5)", "Units, figures, references, limitations stated, MATLAB work another student can rerun"),
]
t2 = doc.add_table(rows=0, cols=2)
t2.style = "Light Grid Accent 1"
for a, b in mk:
    c = t2.add_row().cells
    c[0].text = a
    c[1].text = b
    c[0].paragraphs[0].runs[0].bold = True
doc.add_paragraph()

doc.add_heading("Getting started", level=1)
para("Put the supplied folder somewhere sensible and open MATLAB R2025a or "
     "newer. Read README.md in that folder for a file-by-file description, then "
     "work through the templates in order, filling the cells marked TODO. Each "
     "template tells you the method, gives the governing equations, and names "
     "the Shigley chapter it comes from.")
para("Alongside the four templates your folder holds three supporting "
     "documents (ASSIGNMENT_BRIEF.md, DATA_DICTIONARY.md and "
     "DESIGN_REPORT_TEMPLATE.md), the two license files, and computeShoulderNotchFactors.m, "
     "which returns the shoulder stress concentration factors so you do not "
     "have to read five charts by hand for every candidate geometry. Read one "
     "shoulder off the charts yourself and check the function against it before "
     "you trust it for the rest.")
para("The simulation results are in exports_design_ready, which sits beside the "
     "templates in the same folder. Keep it there and the relative paths in the "
     "templates will resolve. Design_Summary.mat holds the vehicle and gear "
     "data. Each of the three case files holds the design values for that case "
     "in ds_single, one torque-against-time curve you can plot, and a note "
     "saying what that curve is. For the city cycle that curve is the "
     "equivalent sinusoid your Goodman calculation uses, not the raw tracked "
     "torque; README.md explains why and what to watch for.")

doc.add_heading("What a good submission looks like", level=1)
para("The templates contain no starter values. Every value in your design is one you "
     "chose and have to defend. If it passes every check on the first attempt, "
     "you have probably not yet found the check that governs it.")
para("A strong report states which section, which check and which parameter "
     "decided the design, and what you would change first if the car had "
     "more torque.")

# ---------------------------------------------------------------- metadata
now = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
props = doc.core_properties
props.author = AUTHOR
props.last_modified_by = AUTHOR
props.title = TITLE
props.subject = "Mechanical Engineering Design project brief"
props.category = "Course material"
props.comments = ""
props.keywords = ""
props.description = ""
props.created = now
props.modified = now
props.revision = 1

doc.save(OUT)


def strip_tool_metadata(path):
    """Remove the producing-tool fields from docProps/app.xml.

    The python-docx template claims Microsoft Macintosh Word and a 2013 date,
    which is untrue. Naming the real build tool instead is no more useful to a
    student reading File > Info, so these fields are removed altogether and the
    brief carries no tool footprint.
    """
    with zipfile.ZipFile(path) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    for index, (item, data) in enumerate(entries):
        if item.filename == "docProps/app.xml":
            text = data.decode("utf-8")
            for tag in ("Application", "AppVersion", "Company", "Manager", "Template"):
                text = re.sub("<%s>.*?</%s>" % (tag, tag), "", text, flags=re.S)
                text = re.sub("<%s ?/>" % tag, "", text)
            entries[index] = (item, text.encode("utf-8"))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for item, data in entries:
            target.writestr(item, data)


strip_tool_metadata(OUT)

print("written:", OUT, os.path.getsize(OUT), "bytes")
